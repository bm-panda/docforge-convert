"""表格类格式转换：XLSX、XLS、CSV、TSV、ODS。

统一抽象：每份文件视为「多个工作表」，每个工作表为「行列表」。
读取：openpyxl(xlsx) / xlrd(xls) / odfpy(ods) / csv( csv/tsv)
写出：openpyxl(xlsx) / csv / json / markdown / html / reportlab(pdf) / docx / odfpy(ods)
"""

from __future__ import annotations

import csv
import html as html_lib
import json
from pathlib import Path

from docforge.converters.base import ConversionError, register
from docforge.core.utils import read_text_auto

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

try:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    _FONT = "STSong-Light"
except Exception:
    _FONT = "Helvetica"


# ---------------------------------------------------------------------------
# 读取
# ---------------------------------------------------------------------------

def _read_xlsx(path: Path) -> list[tuple[str, list[list]]]:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if c is None else c for c in row])
        sheets.append((ws.title, rows))
    wb.close()
    return sheets


def _read_xls(path: Path) -> list[tuple[str, list[list]]]:
    import xlrd
    wb = xlrd.open_workbook(str(path))
    sheets = []
    for sheet in wb.sheets():
        rows = []
        for r in range(sheet.nrows):
            rows.append([sheet.cell_value(r, c) for c in range(sheet.ncols)])
        sheets.append((sheet.name, rows))
    return sheets


def _read_ods(path: Path) -> list[tuple[str, list[list]]]:
    from odf.opendocument import load as odf_load
    from odf.table import Table as OdTable, TableRow, TableCell
    from odf.text import P
    from odf import teletype
    doc = odf_load(str(path))
    sheets = []
    for table in doc.getElementsByType(OdTable):
        name = table.getAttribute("name") or "Sheet"
        rows = []
        for row in table.getElementsByType(TableRow):
            cells = []
            for cell in row.getElementsByType(TableCell):
                txt = " ".join(teletype.extractText(p) for p in cell.getElementsByType(P))
                cells.append(txt)
            rows.append(cells)
        sheets.append((name, rows))
    return sheets


def _read_delimited(path: Path, delimiter: str) -> list[tuple[str, list[list]]]:
    text = read_text_auto(path)
    reader = csv.reader(text.splitlines(), delimiter=delimiter)
    rows = [row for row in reader]
    return [("Sheet1", rows)]


def read_sheets(src: Path, src_ext: str) -> list[tuple[str, list[list]]]:
    src_ext = src_ext.lower()
    if src_ext == "xlsx":
        return _read_xlsx(src)
    if src_ext == "xls":
        return _read_xls(src)
    if src_ext == "ods":
        return _read_ods(src)
    if src_ext == "tsv":
        return _read_delimited(src, "\t")
    if src_ext == "csv":
        return _read_delimited(src, ",")
    raise ConversionError(f"不支持的表格格式：{src_ext}")


# ---------------------------------------------------------------------------
# 写出
# ---------------------------------------------------------------------------

def _write_csv(rows: list[list], dst: Path, delimiter=","):
    with open(dst, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=delimiter)
        for row in rows:
            w.writerow(["" if c is None else c for c in row])


def _to_md_table(rows: list[list]) -> str:
    if not rows:
        return ""
    ncols = max(len(r) for r in rows)
    norm = [r + [""] * (ncols - len(r)) for r in rows]
    lines = []
    lines.append("| " + " | ".join(str(c) for c in norm[0]) + " |")
    lines.append("| " + " | ".join("---" for _ in range(ncols)) + " |")
    for row in norm[1:]:
        lines.append("| " + " | ".join(str(c) for c in row) + " |")
    return "\n".join(lines)


def _to_html_table(rows: list[list], title: str = "") -> str:
    if not rows:
        return "<p>空表格</p>"
    parts = ["<html><head><meta charset='utf-8'>",
             f"<title>{html_lib.escape(title)}</title></head>",
             "<body style='font-family:SimSun,serif'><table border='1' "
             "cellspacing='0' cellpadding='6' style='border-collapse:collapse'>"]
    for i, row in enumerate(rows):
        tag = "th" if i == 0 else "td"
        parts.append("<tr>" + "".join(f"<{tag}>{html_lib.escape(str(c))}</{tag}>" for c in row) + "</tr>")
    parts.append("</table></body></html>")
    return "".join(parts)


def _to_json(rows: list[list]) -> str:
    if not rows:
        return "[]"
    header = [str(c) for c in rows[0]]
    records = []
    for row in rows[1:]:
        rec = {}
        for i, key in enumerate(header):
            rec[key] = row[i] if i < len(row) else ""
        records.append(rec)
    return json.dumps(records, ensure_ascii=False, indent=2)


def _to_pdf_table(sheets: list[tuple[str, list[list]]], dst: Path, title: str):
    style_par = ParagraphStyle("cell", fontName=_FONT, fontSize=8.5, leading=11)
    header_par = ParagraphStyle("hcell", fontName=_FONT, fontSize=8.5, leading=11)
    flowables = []
    for idx, (name, rows) in enumerate(sheets):
        if not rows:
            continue
        if len(sheets) > 1:
            flowables.append(Paragraph(html_lib.escape(name), ParagraphStyle(
                "sheet", fontName=_FONT, fontSize=13, leading=18, spaceBefore=10)))
        data = []
        for ri, row in enumerate(rows):
            data.append([Paragraph(html_lib.escape(str(c or "")), header_par if ri == 0 else style_par)
                         for c in row])
        tbl = Table(data, repeatRows=1)
        cmds = [
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef1f6")),
        ]
        tbl.setStyle(TableStyle(cmds))
        flowables.append(tbl)
        flowables.append(Spacer(1, 10))
    doc = SimpleDocTemplate(str(dst), pagesize=landscape(A4),
                            leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm,
                            title=title, author="DocForge")
    doc.build(flowables)


def _to_docx(sheets: list[tuple[str, list[list]]], dst: Path, title: str):
    from docx import Document
    from docx.shared import Pt
    from docforge.converters.office_conv import _heading_level  # noqa: F401
    doc = Document()
    for name, rows in sheets:
        if len(sheets) > 1:
            h = doc.add_heading(name, level=1)
            for r in h.runs:
                r.font.name = "Calibri"
        if not rows:
            continue
        ncols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=ncols)
        table.style = "Table Grid"
        for ri, row in enumerate(rows):
            for ci in range(ncols):
                cell = table.cell(ri, ci)
                cell.text = str(row[ci]) if ci < len(row) else ""
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
    doc.save(str(dst))


def _to_ods(sheets: list[tuple[str, list[list]]], dst: Path):
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table as OdTable, TableRow, TableCell
    from odf.text import P
    doc = OpenDocumentSpreadsheet()
    for name, rows in sheets:
        table = OdTable(name=name)
        for row in rows:
            tr = TableRow()
            for cell in row:
                tc = TableCell()
                p = P(text=str(cell) if cell is not None else "")
                tc.addElement(p)
                tr.addElement(tc)
            table.addElement(tr)
        doc.spreadsheet.addElement(table)
    doc.save(str(dst))


# ---------------------------------------------------------------------------
# 通用转换器工厂
# ---------------------------------------------------------------------------

def _mk_sheet_converters():
    """为每种表格源格式批量注册目标转换器。"""

    def factory(src_ext):
        @register(src_ext, "csv")
        def to_csv(src, dst, options=None):
            sheets = read_sheets(src, src_ext)
            _write_csv(sheets[0][1], dst, ",")

        @register(src_ext, "tsv")
        def to_tsv(src, dst, options=None):
            sheets = read_sheets(src, src_ext)
            _write_csv(sheets[0][1], dst, "\t")

        @register(src_ext, "json")
        def to_json(src, dst, options=None):
            sheets = read_sheets(src, src_ext)
            dst.write_text(_to_json(sheets[0][1]), encoding="utf-8")

        @register(src_ext, "md")
        def to_md(src, dst, options=None):
            sheets = read_sheets(src, src_ext)
            parts = []
            for name, rows in sheets:
                if len(sheets) > 1:
                    parts.append(f"## {name}\n")
                parts.append(_to_md_table(rows))
            dst.write_text("\n\n".join(parts), encoding="utf-8")

        @register(src_ext, "html")
        def to_html(src, dst, options=None):
            sheets = read_sheets(src, src_ext)
            parts = []
            for name, rows in sheets:
                if len(sheets) > 1:
                    parts.append(f"<h2>{html_lib.escape(name)}</h2>")
                parts.append(_to_html_table(rows, name))
            dst.write_text(
                "<html><head><meta charset='utf-8'><title>表格文档</title></head>"
                f"<body>{''.join(parts)}</body></html>", encoding="utf-8")

        @register(src_ext, "pdf")
        def to_pdf(src, dst, options=None):
            # Office 表格：本机装了 MS Excel 时优先走 COM 高保真导出，失败回退纯 Python 表格渲染
            if src_ext in ("xlsx", "xls", "ods"):
                try:
                    from docforge.converters.office_com import excel_to_pdf
                    if excel_to_pdf(src, dst):
                        return
                except Exception:
                    pass
            sheets = read_sheets(src, src_ext)
            _to_pdf_table(sheets, dst, src.stem)

        @register(src_ext, "docx")
        def to_docx(src, dst, options=None):
            sheets = read_sheets(src, src_ext)
            _to_docx(sheets, dst, src.stem)

        @register(src_ext, "xlsx")
        def to_xlsx(src, dst, options=None):
            if src_ext == "xlsx":
                # 直接复制文件（含全部样式）
                import shutil
                shutil.copyfile(src, dst)
                return
            sheets = read_sheets(src, src_ext)
            import openpyxl
            wb = openpyxl.Workbook()
            wb.remove(wb.active)
            for name, rows in sheets:
                ws = wb.create_sheet(title=(name or "Sheet")[:31])
                for row in rows:
                    ws.append([c for c in row])
            wb.save(str(dst))

        if src_ext != "ods":
            @register(src_ext, "ods")
            def to_ods(src, dst, options=None):
                sheets = read_sheets(src, src_ext)
                _to_ods(sheets, dst)

    for ext in ("xlsx", "xls", "csv", "tsv", "ods"):
        factory(ext)


_mk_sheet_converters()


# ---------------------------------------------------------------------------
# CSV → TXT 等补充转换
# ---------------------------------------------------------------------------

@register("csv", "txt")
def csv_to_txt(src: Path, dst: Path, options=None):
    sheets = read_sheets(src, "csv")
    dst.write_text(_to_md_table(sheets[0][1]), encoding="utf-8")


@register("tsv", "txt")
def tsv_to_txt(src: Path, dst: Path, options=None):
    sheets = read_sheets(src, "tsv")
    dst.write_text(_to_md_table(sheets[0][1]), encoding="utf-8")
