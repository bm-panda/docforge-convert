"""Office 文档类转换：DOCX、PPTX、ODT（纯 Python 解析）。

DOCX/PPTX/ODT 均可在无 Office 环境下读取文本并输出各类格式；
DOCX→PDF 优先尝试 Word COM（见 office_com），失败回退纯文本渲染。
"""

from __future__ import annotations

import html as html_lib
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from docforge.converters.base import ConversionError, register
from docforge.converters.html_pdf import text_to_pdf
from docforge.converters.text_conv import _html_to_docx, _text_to_docx, _text_to_md
from docforge.core.utils import read_text_auto

try:
    from odf import teletype
    from odf.opendocument import load as odf_load
    from odf.text import H, P
    _ODF_OK = True
except Exception:
    _ODF_OK = False


# ---------------------------------------------------------------------------
# DOCX 读取
# ---------------------------------------------------------------------------

def iter_block_items(parent):
    """按文档顺序迭代段落与表格。"""
    parent_elm = parent.element.body
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _heading_level(paragraph: Paragraph) -> int:
    name = (paragraph.style.name or "").lower()
    if name.startswith("heading"):
        try:
            return int(name.split()[-1])
        except ValueError:
            return 1
    if name.startswith("标题"):
        try:
            return int(name.replace("标题", "").strip() or 1)
        except ValueError:
            return 1
    return 0


def _paragraph_images(paragraph: Paragraph, document_part) -> list[str]:
    """返回该段落内所有内嵌图片的 data URI 列表。"""
    from docx.oxml.ns import qn as _qn
    import base64
    uris = []
    for blip in paragraph._element.iter(_qn("a:blip")):
        r_id = blip.get(_qn("r:embed"))
        if not r_id:
            continue
        try:
            part = document_part.related_parts[r_id]
        except (KeyError, AttributeError):
            continue
        data = getattr(part, "blob", None)
        if not data:
            continue
        mime = getattr(part, "content_type", None) or "image/png"
        uris.append(f"data:{mime};base64," + base64.b64encode(data).decode("ascii"))
    return uris


def _docx_extract(src: Path) -> list[tuple[int, str, list]]:
    """返回 [(level, text, image_uris)]，level=0 表示正文段落。"""
    doc = Document(str(src))
    out = []
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            images = _paragraph_images(block, doc.part)
            if text:
                out.append((_heading_level(block), text, images))
            else:
                for uri in images:
                    out.append((0, "", [uri]))
        else:
            for row in block.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                out.append((0, " | ".join(cells), []))
    return out


def _docx_to_text(src: Path) -> str:
    return "\n".join(t for _, t, _ in _docx_extract(src))


def _docx_to_md(src: Path) -> str:
    lines = []
    for level, text, images in _docx_extract(src):
        if images:
            for uri in images:
                lines.append(f"![图片]({uri})")
        if text:
            lines.append(f"{'#' * min(level, 6)} {text}" if level else text)
    return "\n\n".join(lines)


def _docx_to_html(src: Path) -> str:
    parts = ["<html><head><meta charset='utf-8'>",
             f"<title>{html_lib.escape(src.stem)}</title></head>",
             "<body style='font-family:SimSun,serif;line-height:1.7'>"]
    for level, text, images in _docx_extract(src):
        if images:
            for uri in images:
                parts.append(f'<p><img src="{uri}" style="max-width:100%"/></p>')
        if text:
            if level == 0:
                parts.append(f"<p>{html_lib.escape(text)}</p>")
            else:
                tag = f"h{min(level, 6)}"
                parts.append(f"<{tag}>{html_lib.escape(text)}</{tag}>")
    parts.append("</body></html>")
    return "".join(parts)


# ---------------------------------------------------------------------------
# DOCX 目标转换
# ---------------------------------------------------------------------------

@register("docx", "txt")
def docx_to_txt(src: Path, dst: Path, options=None):
    dst.write_text(_docx_to_text(src), encoding="utf-8")


@register("docx", "md")
def docx_to_md(src: Path, dst: Path, options=None):
    dst.write_text(_docx_to_md(src), encoding="utf-8")


@register("docx", "html")
def docx_to_html(src: Path, dst: Path, options=None):
    dst.write_text(_docx_to_html(src), encoding="utf-8")


@register("docx", "pdf")
def docx_to_pdf(src: Path, dst: Path, options=None):
    # 优先走 Word COM 高保真转换
    try:
        from docforge.converters.office_com import word_to_pdf
        if word_to_pdf(src, dst):
            return
    except Exception:
        pass
    text = _docx_to_text(src)
    if not text.strip():
        raise ConversionError("DOCX 无可提取文本，无法渲染 PDF")
    text_to_pdf(text, dst, title=src.stem)


@register("docx", "rtf")
def docx_to_rtf(src: Path, dst: Path, options=None):
    try:
        from docforge.converters.office_com import word_save_as
        if word_save_as(src, dst, fmt=6):  # wdFormatRTF = 6
            return
    except Exception:
        pass
    from docforge.converters.text_conv import _text_to_rtf
    dst.write_text(_text_to_rtf(_docx_to_text(src)), encoding="utf-8")


# ---------------------------------------------------------------------------
# PPTX 读取
# ---------------------------------------------------------------------------

def _pptx_extract(src: Path) -> list[tuple[int, str]]:
    from pptx import Presentation
    prs = Presentation(str(src))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        out.append((0, text))
            if shape.shape_type == 19:  # table
                try:
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells]
                        out.append((0, " | ".join(cells)))
                except Exception:
                    pass
    return out


def _pptx_to_text(src: Path) -> str:
    return "\n".join(t for _, t in _pptx_extract(src))


def _pptx_to_md(src: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(src))
    lines = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in para.runs).strip()
                    if t:
                        texts.append(t)
        if texts:
            lines.append(f"## 幻灯片 {i}")
            for t in texts:
                lines.append("- " + t)
            lines.append("")
    return "\n".join(lines)


@register("pptx", "txt")
def pptx_to_txt(src: Path, dst: Path, options=None):
    dst.write_text(_pptx_to_text(src), encoding="utf-8")


@register("pptx", "md")
def pptx_to_md(src: Path, dst: Path, options=None):
    dst.write_text(_pptx_to_md(src), encoding="utf-8")


@register("pptx", "html")
def pptx_to_html(src: Path, dst: Path, options=None):
    lines = []
    for level, text in _pptx_extract(src):
        if level == 0:
            lines.append(f"<p>{html_lib.escape(text)}</p>")
    dst.write_text(
        "<html><head><meta charset='utf-8'></head>"
        f"<body style='font-family:SimSun,serif;line-height:1.7'>{''.join(lines)}</body></html>",
        encoding="utf-8",
    )


@register("pptx", "pdf")
def pptx_to_pdf(src: Path, dst: Path, options=None):
    # 优先 PowerPoint COM 高保真导出
    try:
        from docforge.converters.office_com import ppt_to_pdf
        if ppt_to_pdf(src, dst):
            return
    except Exception:
        pass
    text = _pptx_to_text(src)
    if not text.strip():
        raise ConversionError("PPTX 无可提取文本，无法渲染 PDF")
    text_to_pdf(text, dst, title=src.stem)


# ---------------------------------------------------------------------------
# ODT 读取与转换
# ---------------------------------------------------------------------------

def _odt_extract(src: Path) -> list[tuple[int, str]]:
    if not _ODF_OK:
        raise ConversionError("缺少 odfpy 库，无法读取 ODT")
    doc = odf_load(str(src))
    out = []
    for el in doc.getElementsByType(P):
        text = teletype.extractText(el).strip()
        if text:
            out.append((0, text))
    for el in doc.getElementsByType(H):
        text = teletype.extractText(el).strip()
        if text:
            try:
                lvl = int(el.getAttribute("outlinelevel") or 1)
            except (TypeError, ValueError):
                lvl = 1
            out.append((min(lvl, 6), text))
    return out


def _odt_to_text(src: Path) -> str:
    return "\n".join(t for _, t in _odt_extract(src))


@register("odt", "txt")
def odt_to_txt(src: Path, dst: Path, options=None):
    dst.write_text(_odt_to_text(src), encoding="utf-8")


@register("odt", "docx")
def odt_to_docx(src: Path, dst: Path, options=None):
    blocks = _odt_extract(src)
    html_parts = []
    for level, text in blocks:
        if level:
            html_parts.append(f"<h{level}>{html_lib.escape(text)}</h{level}>")
        else:
            html_parts.append(f"<p>{html_lib.escape(text)}</p>")
    _html_to_docx("".join(html_parts), dst)


@register("odt", "md")
def odt_to_md(src: Path, dst: Path, options=None):
    lines = []
    for level, text in _odt_extract(src):
        lines.append(f"{'#' * min(level, 6)} {text}" if level else text)
    dst.write_text("\n\n".join(lines), encoding="utf-8")


@register("odt", "html")
def odt_to_html(src: Path, dst: Path, options=None):
    parts = ["<html><head><meta charset='utf-8'></head><body>"]
    for level, text in _odt_extract(src):
        tag = f"h{min(level, 6)}" if level else "p"
        parts.append(f"<{tag}>{html_lib.escape(text)}</{tag}>")
    parts.append("</body></html>")
    dst.write_text("".join(parts), encoding="utf-8")


@register("odt", "pdf")
def odt_to_pdf(src: Path, dst: Path, options=None):
    text = _odt_to_text(src)
    if not text.strip():
        raise ConversionError("ODT 无可提取文本，无法渲染 PDF")
    text_to_pdf(text, dst, title=src.stem)
